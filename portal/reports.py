import os
import io
from datetime import datetime
from decimal import Decimal
from django.conf import settings
from django.utils import timezone
from django.db.models import Avg, Count, Max, Min, Q
from django.http import HttpResponse

import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib import colors
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle, HRFlowable, KeepTogether, Image
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle

from .models import AcademicClass, Student, Exam, ExamAttempt


# ==============================================================================
# DATA AGGREGATION HELPERS
# ==============================================================================

def get_class_report_data(class_id):
    """
    Gathers detailed class-wise metrics:
    - How many exams scheduled
    - How many students enrolled
    - Per-exam: enrolled, attempted count, pending count, total submissions, pass/fail, avg %, highest score
    - Per-student: assigned tests, attempted tests, pending tests, avg %, pass count
    """
    academic_class = AcademicClass.objects.get(id=class_id)
    students = academic_class.students.all().order_by('roll_number', 'name')
    total_students = students.count()

    exams = academic_class.exams.all().order_by('-created_at')
    total_exams_scheduled = exams.count()

    class_attempts = ExamAttempt.objects.filter(
        student__student_class=academic_class,
        exam__in=exams
    )
    total_attempts_count = class_attempts.count()

    class_avg_pct = class_attempts.aggregate(Avg('percentage'))['percentage__avg'] or 0
    passed_attempts_count = class_attempts.filter(is_passed=True).count()
    failed_attempts_count = class_attempts.filter(is_passed=False, is_completed=True).count()
    overall_pass_rate = (passed_attempts_count / total_attempts_count * 100) if total_attempts_count > 0 else 0

    # 1. Per-Exam Breakdown
    exam_stats = []
    for exam in exams:
        attempts_for_exam = class_attempts.filter(exam=exam)
        distinct_attempted_students = attempts_for_exam.values('student').distinct().count()
        pending_students_count = max(0, total_students - distinct_attempted_students)

        exam_total_submissions = attempts_for_exam.count()
        exam_passed = attempts_for_exam.filter(is_passed=True).count()
        exam_failed = attempts_for_exam.filter(is_passed=False, is_completed=True).count()
        exam_avg_pct = attempts_for_exam.aggregate(Avg('percentage'))['percentage__avg'] or 0
        exam_max_score = attempts_for_exam.aggregate(Max('score_obtained'))['score_obtained__max'] or 0

        exam_stats.append({
            'exam': exam,
            'title': exam.title,
            'subject': exam.subject or 'General',
            'duration_minutes': exam.duration_minutes,
            'total_marks': exam.total_marks,
            'passing_percentage': exam.passing_percentage,
            'questions_count': exam.total_questions_count,
            'enrolled_students': total_students,
            'attempted_students': distinct_attempted_students,
            'pending_students': pending_students_count,
            'total_submissions': exam_total_submissions,
            'passed_count': exam_passed,
            'failed_count': exam_failed,
            'avg_percentage': round(float(exam_avg_pct), 1),
            'highest_score': round(float(exam_max_score), 2),
            'participation_rate': round((distinct_attempted_students / total_students * 100), 1) if total_students > 0 else 0
        })

    # 2. Per-Student Breakdown
    student_stats = []
    for s in students:
        s_attempts = class_attempts.filter(student=s)
        distinct_exams_attempted = s_attempts.values('exam').distinct().count()
        pending_exams_count = max(0, total_exams_scheduled - distinct_exams_attempted)
        total_s_attempts = s_attempts.count()

        s_avg_pct = s_attempts.aggregate(Avg('percentage'))['percentage__avg'] or 0
        s_best_pct = s_attempts.aggregate(Max('percentage'))['percentage__max'] or 0
        s_passed = s_attempts.filter(is_passed=True).count()

        if distinct_exams_attempted == 0:
            status = 'Not Started'
        elif pending_exams_count == 0:
            status = 'Completed All'
        else:
            status = 'In Progress'

        student_stats.append({
            'student': s,
            'name': s.name,
            'roll_number': s.roll_number,
            'phone': s.phone or '-',
            'assigned_exams': total_exams_scheduled,
            'attempted_exams': distinct_exams_attempted,
            'pending_exams': pending_exams_count,
            'total_attempts': total_s_attempts,
            'passed_count': s_passed,
            'avg_percentage': round(float(s_avg_pct), 1),
            'best_percentage': round(float(s_best_pct), 1),
            'status': status
        })

    return {
        'academic_class': academic_class,
        'class_name': academic_class.name,
        'class_code': academic_class.code or 'BATCH',
        'total_students': total_students,
        'total_exams_scheduled': total_exams_scheduled,
        'total_attempts_count': total_attempts_count,
        'class_avg_pct': round(float(class_avg_pct), 1),
        'passed_attempts_count': passed_attempts_count,
        'failed_attempts_count': failed_attempts_count,
        'overall_pass_rate': round(float(overall_pass_rate), 1),
        'exam_stats': exam_stats,
        'student_stats': student_stats,
        'generated_at': timezone.now(),
    }


def get_student_report_data(student_id):
    """
    Gathers detailed student-wise metrics:
    - Total tests assigned in class
    - Total tests attempted vs pending
    - Complete chronological attempt history with marks, %, result
    - Pending tests not yet attempted
    """
    student = Student.objects.select_related('student_class').get(id=student_id)
    academic_class = student.student_class

    assigned_exams = academic_class.exams.all().order_by('-created_at')
    total_assigned_exams = assigned_exams.count()

    attempts = student.attempts.select_related('exam').order_by('-started_at')
    total_attempts_count = attempts.count()

    distinct_attempted_exam_ids = set(attempts.values_list('exam_id', flat=True))
    attempted_exams_count = len(distinct_attempted_exam_ids)
    pending_exams_count = max(0, total_assigned_exams - attempted_exams_count)

    overall_avg_pct = attempts.aggregate(Avg('percentage'))['percentage__avg'] or 0
    best_pct = attempts.aggregate(Max('percentage'))['percentage__max'] or 0
    passed_count = attempts.filter(is_passed=True).count()
    failed_count = attempts.filter(is_passed=False, is_completed=True).count()

    # Pending exams list
    pending_exams = [e for e in assigned_exams if e.id not in distinct_attempted_exam_ids]

    attempt_rows = []
    for att in attempts:
        attempt_rows.append({
            'attempt': att,
            'id': att.id,
            'attempt_number': att.attempt_number,
            'exam_title': att.exam.title,
            'subject': att.exam.subject or 'General',
            'started_at': att.started_at,
            'time_taken_minutes': round(att.time_taken_seconds / 60, 1) if att.time_taken_seconds else 0,
            'score_obtained': att.score_obtained,
            'max_score': att.max_score,
            'percentage': round(float(att.percentage), 1),
            'is_passed': att.is_passed,
        })

    return {
        'student': student,
        'name': student.name,
        'roll_number': student.roll_number,
        'phone': student.phone or '-',
        'email': student.email or '-',
        'academic_class': academic_class,
        'class_name': academic_class.name,
        'total_assigned_exams': total_assigned_exams,
        'attempted_exams_count': attempted_exams_count,
        'pending_exams_count': pending_exams_count,
        'total_attempts_count': total_attempts_count,
        'overall_avg_pct': round(float(overall_avg_pct), 1),
        'best_pct': round(float(best_pct), 1),
        'passed_count': passed_count,
        'failed_count': failed_count,
        'attempt_rows': attempt_rows,
        'pending_exams': pending_exams,
        'generated_at': timezone.now(),
    }


# ==============================================================================
# EXCEL GENERATION (.xlsx using openpyxl)
# ==============================================================================

def generate_class_excel(data):
    """
    Creates a styled Microsoft Excel workbook (.xlsx) for a Class Report.
    """
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Class Performance Summary"
    ws.views.sheetView[0].showGridLines = True

    # Color definitions
    navy_fill = PatternFill(start_color="1E3A8A", end_color="1E3A8A", fill_type="solid")
    accent_fill = PatternFill(start_color="F59E0B", end_color="F59E0B", fill_type="solid")
    light_fill = PatternFill(start_color="F8FAFC", end_color="F8FAFC", fill_type="solid")
    pass_fill = PatternFill(start_color="DCFCE7", end_color="DCFCE7", fill_type="solid")
    fail_fill = PatternFill(start_color="FEE2E2", end_color="FEE2E2", fill_type="solid")
    header_fill = PatternFill(start_color="0F172A", end_color="0F172A", fill_type="solid")

    font_title = Font(name="Calibri", size=16, bold=True, color="1E3A8A")
    font_subtitle = Font(name="Calibri", size=11, bold=True, color="475569")
    font_timestamp = Font(name="Calibri", size=9, italic=True, color="64748B")
    font_header = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    font_bold = Font(name="Calibri", size=11, bold=True, color="0F172A")
    font_regular = Font(name="Calibri", size=10, color="1E293B")
    font_pass = Font(name="Calibri", size=10, bold=True, color="15803D")
    font_fail = Font(name="Calibri", size=10, bold=True, color="B91C1C")

    thin_border = Border(
        left=Side(style='thin', color='CBD5E1'),
        right=Side(style='thin', color='CBD5E1'),
        top=Side(style='thin', color='CBD5E1'),
        bottom=Side(style='thin', color='CBD5E1')
    )

    # 1. School Letterhead Banner
    ws.append(["Arohan Academy English School"])
    ws.cell(row=1, column=1).font = font_title
    ws.append([f"Class Examination Report: {data['class_name']} ({data['class_code']})"])
    ws.cell(row=2, column=1).font = font_subtitle
    ws.append([f"Generated: {data['generated_at'].strftime('%d %B %Y, %I:%M %p')} | Confidential Academic Record"])
    ws.cell(row=3, column=1).font = font_timestamp
    ws.append([])  # Blank row 4

    # 2. Executive KPI Metrics Row
    ws.append(["EXECUTIVE KPI SUMMARY"])
    ws.cell(row=5, column=1).font = font_bold

    kpi_headers = ["Total Enrolled Students", "Scheduled Mock Tests", "Total Attempt Submissions", "Class Average Score", "Overall Pass Rate"]
    kpi_values = [
        data['total_students'],
        data['total_exams_scheduled'],
        data['total_attempts_count'],
        f"{data['class_avg_pct']}%",
        f"{data['overall_pass_rate']}%"
    ]
    ws.append(kpi_headers)
    ws.append(kpi_values)

    for col in range(1, len(kpi_headers) + 1):
        cell_h = ws.cell(row=6, column=col)
        cell_h.font = font_header
        cell_h.fill = navy_fill
        cell_h.alignment = Alignment(horizontal='center', vertical='center')
        cell_h.border = thin_border

        cell_v = ws.cell(row=7, column=col)
        cell_v.font = font_bold
        cell_v.alignment = Alignment(horizontal='center', vertical='center')
        cell_v.fill = light_fill
        cell_v.border = thin_border

    ws.append([])  # Blank row 8

    # 3. Table 1: Exam Schedule & Attempt Status Breakdown
    ws.append(["1. MOCK TESTS SCHEDULE & STUDENT PARTICIPATION BREAKDOWN"])
    ws.cell(row=9, column=1).font = font_bold

    exam_headers = [
        "Mock Test Title", "Subject", "Total Marks", "Pass %",
        "Enrolled", "Attempted", "Pending", "Submissions", "Passed", "Failed", "Avg %", "Highest Score"
    ]
    ws.append(exam_headers)
    for col in range(1, len(exam_headers) + 1):
        cell = ws.cell(row=10, column=col)
        cell.font = font_header
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal='center', vertical='center')
        cell.border = thin_border

    curr_row = 11
    for ex in data['exam_stats']:
        row_data = [
            ex['title'],
            ex['subject'],
            ex['total_marks'],
            f"{ex['passing_percentage']}%",
            ex['enrolled_students'],
            ex['attempted_students'],
            ex['pending_students'],
            ex['total_submissions'],
            ex['passed_count'],
            ex['failed_count'],
            f"{ex['avg_percentage']}%",
            ex['highest_score']
        ]
        ws.append(row_data)
        for col in range(1, len(row_data) + 1):
            cell = ws.cell(row=curr_row, column=col)
            cell.font = font_regular
            cell.border = thin_border
            if col == 1:
                cell.alignment = Alignment(horizontal='left')
            else:
                cell.alignment = Alignment(horizontal='center')
            if col == 7 and ex['pending_students'] > 0:
                cell.fill = PatternFill(start_color="FEF3C7", end_color="FEF3C7", fill_type="solid")  # Amber for pending
        curr_row += 1

    ws.append([])  # Blank row
    curr_row += 1

    # 4. Table 2: Student-wise Participation & Performance Roster
    ws.append(["2. ENROLLED STUDENTS PARTICIPATION & PROGRESS ROSTER"])
    ws.cell(row=curr_row, column=1).font = font_bold
    curr_row += 1

    st_headers = [
        "Roll Number", "Student Name", "Contact",
        "Assigned Tests", "Attempted Tests", "Pending Tests",
        "Submissions", "Passed Count", "Average %", "Best %", "Completion Status"
    ]
    ws.append(st_headers)
    for col in range(1, len(st_headers) + 1):
        cell = ws.cell(row=curr_row, column=col)
        cell.font = font_header
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal='center', vertical='center')
        cell.border = thin_border
    curr_row += 1

    for st in data['student_stats']:
        row_data = [
            st['roll_number'],
            st['name'],
            st['phone'],
            st['assigned_exams'],
            st['attempted_exams'],
            st['pending_exams'],
            st['total_attempts'],
            st['passed_count'],
            f"{st['avg_percentage']}%",
            f"{st['best_percentage']}%",
            st['status']
        ]
        ws.append(row_data)
        for col in range(1, len(row_data) + 1):
            cell = ws.cell(row=curr_row, column=col)
            cell.font = font_regular
            cell.border = thin_border
            if col == 2:
                cell.alignment = Alignment(horizontal='left')
            else:
                cell.alignment = Alignment(horizontal='center')

            if col == 11:
                if st['status'] == 'Completed All':
                    cell.fill = pass_fill
                    cell.font = font_pass
                elif st['status'] == 'Not Started':
                    cell.fill = fail_fill
                    cell.font = font_fail
        curr_row += 1

    # Auto-fit column widths
    for col in ws.columns:
        col_letter = get_column_letter(col[0].column)
        max_len = 0
        for cell in col:
            val_str = str(cell.value or '')
            if cell.row in [1, 2, 3]:
                continue
            if len(val_str) > max_len:
                max_len = len(val_str)
        ws.column_dimensions[col_letter].width = max(max_len + 3, 12)

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf.getvalue()


def generate_student_excel(data):
    """
    Creates a styled Microsoft Excel workbook (.xlsx) for an individual Student Report.
    """
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Student Examination Report"
    ws.views.sheetView[0].showGridLines = True

    navy_fill = PatternFill(start_color="1E3A8A", end_color="1E3A8A", fill_type="solid")
    light_fill = PatternFill(start_color="F8FAFC", end_color="F8FAFC", fill_type="solid")
    pass_fill = PatternFill(start_color="DCFCE7", end_color="DCFCE7", fill_type="solid")
    fail_fill = PatternFill(start_color="FEE2E2", end_color="FEE2E2", fill_type="solid")
    header_fill = PatternFill(start_color="0F172A", end_color="0F172A", fill_type="solid")

    font_title = Font(name="Calibri", size=16, bold=True, color="1E3A8A")
    font_subtitle = Font(name="Calibri", size=11, bold=True, color="475569")
    font_timestamp = Font(name="Calibri", size=9, italic=True, color="64748B")
    font_header = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    font_bold = Font(name="Calibri", size=11, bold=True, color="0F172A")
    font_regular = Font(name="Calibri", size=10, color="1E293B")
    font_pass = Font(name="Calibri", size=10, bold=True, color="15803D")
    font_fail = Font(name="Calibri", size=10, bold=True, color="B91C1C")

    thin_border = Border(
        left=Side(style='thin', color='CBD5E1'),
        right=Side(style='thin', color='CBD5E1'),
        top=Side(style='thin', color='CBD5E1'),
        bottom=Side(style='thin', color='CBD5E1')
    )

    ws.append(["Arohan Academy English School - Student Performance Audit"])
    ws.cell(row=1, column=1).font = font_title
    ws.append([f"Student: {data['name']} (Roll: {data['roll_number']}) | Class: {data['class_name']}"])
    ws.cell(row=2, column=1).font = font_subtitle
    ws.append([f"Generated: {data['generated_at'].strftime('%d %B %Y, %I:%M %p')} | Contact: {data['phone']}"])
    ws.cell(row=3, column=1).font = font_timestamp
    ws.append([])

    # KPI Metrics
    ws.append(["EXAMINATION SUMMARY OVERVIEW"])
    ws.cell(row=5, column=1).font = font_bold

    kpi_headers = ["Assigned Tests", "Attempted Tests", "Pending Tests", "Total Submissions", "Overall Avg %", "Best Score %"]
    kpi_values = [
        data['total_assigned_exams'],
        data['attempted_exams_count'],
        data['pending_exams_count'],
        data['total_attempts_count'],
        f"{data['overall_avg_pct']}%",
        f"{data['best_pct']}%"
    ]
    ws.append(kpi_headers)
    ws.append(kpi_values)

    for col in range(1, len(kpi_headers) + 1):
        cell_h = ws.cell(row=6, column=col)
        cell_h.font = font_header
        cell_h.fill = navy_fill
        cell_h.alignment = Alignment(horizontal='center', vertical='center')
        cell_h.border = thin_border

        cell_v = ws.cell(row=7, column=col)
        cell_v.font = font_bold
        cell_v.alignment = Alignment(horizontal='center', vertical='center')
        cell_v.fill = light_fill
        cell_v.border = thin_border

    ws.append([])

    # Table 1: Complete Attempt History
    ws.append(["1. CHRONOLOGICAL MOCK TEST ATTEMPTS LEDGER"])
    ws.cell(row=9, column=1).font = font_bold

    att_headers = [
        "Attempt #", "Mock Test Title", "Subject", "Date & Time",
        "Time Taken (min)", "Score Obtained", "Maximum Marks", "Percentage (%)", "Result Status"
    ]
    ws.append(att_headers)
    for col in range(1, len(att_headers) + 1):
        cell = ws.cell(row=10, column=col)
        cell.font = font_header
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal='center', vertical='center')
        cell.border = thin_border

    curr_row = 11
    if data['attempt_rows']:
        for row in data['attempt_rows']:
            status_text = "PASSED" if row['is_passed'] else "FAILED"
            row_data = [
                f"#{row['attempt_number']}",
                row['exam_title'],
                row['subject'],
                row['started_at'].strftime('%d %b %Y, %I:%M %p'),
                row['time_taken_minutes'],
                row['score_obtained'],
                row['max_score'],
                f"{row['percentage']}%",
                status_text
            ]
            ws.append(row_data)
            for col in range(1, len(row_data) + 1):
                cell = ws.cell(row=curr_row, column=col)
                cell.font = font_regular
                cell.border = thin_border
                cell.alignment = Alignment(horizontal='center')
                if col == 2:
                    cell.alignment = Alignment(horizontal='left')
                if col == 9:
                    if row['is_passed']:
                        cell.fill = pass_fill
                        cell.font = font_pass
                    else:
                        cell.fill = fail_fill
                        cell.font = font_fail
            curr_row += 1
    else:
        ws.append(["No mock test attempts recorded yet for this student."])
        ws.cell(row=curr_row, column=1).font = font_timestamp
        curr_row += 1

    ws.append([])
    curr_row += 1

    # Table 2: Pending Tests
    ws.append(["2. PENDING MOCK TESTS (YET TO ATTEMPT)"])
    ws.cell(row=curr_row, column=1).font = font_bold
    curr_row += 1

    pending_headers = ["Mock Test Title", "Subject", "Duration (min)", "Total Marks", "Passing %", "Status"]
    ws.append(pending_headers)
    for col in range(1, len(pending_headers) + 1):
        cell = ws.cell(row=curr_row, column=col)
        cell.font = font_header
        cell.fill = PatternFill(start_color="B45309", end_color="B45309", fill_type="solid")
        cell.alignment = Alignment(horizontal='center', vertical='center')
        cell.border = thin_border
    curr_row += 1

    if data['pending_exams']:
        for p_exam in data['pending_exams']:
            row_data = [
                p_exam.title,
                p_exam.subject or 'General',
                p_exam.duration_minutes,
                p_exam.total_marks,
                f"{p_exam.passing_percentage}%",
                "PENDING"
            ]
            ws.append(row_data)
            for col in range(1, len(row_data) + 1):
                cell = ws.cell(row=curr_row, column=col)
                cell.font = font_regular
                cell.border = thin_border
                cell.alignment = Alignment(horizontal='center')
                if col == 1:
                    cell.alignment = Alignment(horizontal='left')
                if col == 6:
                    cell.fill = PatternFill(start_color="FEF3C7", end_color="FEF3C7", fill_type="solid")
            curr_row += 1
    else:
        ws.append(["All assigned mock tests have been attempted! No pending tests."])
        ws.cell(row=curr_row, column=1).font = font_pass
        curr_row += 1

    for col in ws.columns:
        col_letter = get_column_letter(col[0].column)
        max_len = 0
        for cell in col:
            val_str = str(cell.value or '')
            if cell.row in [1, 2, 3]:
                continue
            if len(val_str) > max_len:
                max_len = len(val_str)
        ws.column_dimensions[col_letter].width = max(max_len + 3, 12)

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf.getvalue()


# ==============================================================================
# REPORTLAB PDF GENERATION
# ==============================================================================

def get_report_logo_flowable(logo_size=50):
    """
    Returns an Image Flowable for the school logo if the image exists,
    or None if not found.
    """
    possible_paths = [
        os.path.join(settings.BASE_DIR, 'static', 'images', 'logo.png'),
        os.path.join(settings.BASE_DIR, 'staticfiles', 'images', 'logo.png'),
    ]
    for p in possible_paths:
        if os.path.exists(p):
            try:
                img = Image(p, width=logo_size, height=logo_size)
                return img
            except Exception:
                pass
    return None


def create_report_header_flowable(title_paragraphs, logo_size=50, total_width=770):
    """
    Creates a 2-column Table flowable with the school logo on the left
    and the title details on the right. Falls back gracefully to paragraphs if logo is unavailable.
    """
    logo_img = get_report_logo_flowable(logo_size=logo_size)
    if logo_img:
        logo_col_width = logo_size + 14
        text_col_width = max(200, total_width - logo_col_width)
        header_table = Table([[logo_img, title_paragraphs]], colWidths=[logo_col_width, text_col_width])
        header_table.setStyle(TableStyle([
            ('VALIGN', (0, 0), (-1, -1), 'MIDDLE'),
            ('LEFTPADDING', (0, 0), (-1, -1), 0),
            ('RIGHTPADDING', (0, 0), (-1, -1), 0),
            ('TOPPADDING', (0, 0), (-1, -1), 0),
            ('BOTTOMPADDING', (0, 0), (-1, -1), 2),
            ('RIGHTPADDING', (0, 0), (0, 0), 12),
        ]))
        return header_table
    return title_paragraphs


def generate_class_pdf(data):
    """
    Generates a print-ready Landscape A4 PDF document for Class-wise performance using ReportLab.
    """
    buf = io.BytesIO()
    doc = SimpleDocTemplate(
        buf,
        pagesize=landscape(A4),
        leftMargin=36,
        rightMargin=36,
        topMargin=28,
        bottomMargin=28
    )

    styles = getSampleStyleSheet()

    # Custom styles
    title_style = ParagraphStyle(
        'DocTitle',
        parent=styles['Heading1'],
        fontName='Helvetica-Bold',
        fontSize=15,
        leading=18,
        textColor=colors.HexColor('#1E3A8A'),
        alignment=0,
        spaceAfter=2
    )

    subtitle_style = ParagraphStyle(
        'DocSubtitle',
        parent=styles['Normal'],
        fontName='Helvetica-Bold',
        fontSize=10,
        leading=13,
        textColor=colors.HexColor('#D97706'),
        spaceAfter=3
    )

    meta_style = ParagraphStyle(
        'DocMeta',
        parent=styles['Normal'],
        fontName='Helvetica',
        fontSize=8,
        leading=10,
        textColor=colors.HexColor('#64748B'),
        spaceAfter=0
    )

    heading_section = ParagraphStyle(
        'SectionHeading',
        parent=styles['Heading2'],
        fontName='Helvetica-Bold',
        fontSize=11,
        leading=14,
        textColor=colors.HexColor('#0F172A'),
        spaceBefore=8,
        spaceAfter=6
    )

    cell_style = ParagraphStyle(
        'CellText',
        parent=styles['Normal'],
        fontName='Helvetica',
        fontSize=8.5,
        leading=11,
        textColor=colors.HexColor('#1E293B')
    )

    cell_bold = ParagraphStyle(
        'CellBold',
        parent=styles['Normal'],
        fontName='Helvetica-Bold',
        fontSize=8.5,
        leading=11,
        textColor=colors.HexColor('#0F172A'),
        alignment=1
    )

    cell_header = ParagraphStyle(
        'CellHeader',
        parent=styles['Normal'],
        fontName='Helvetica-Bold',
        fontSize=8.5,
        leading=11,
        textColor=colors.white,
        alignment=1
    )

    elements = []

    # 1. School Header Banner with Logo on Left
    title_paragraphs = [
        Paragraph("Arohan Academy English School", title_style),
        Paragraph(f"Class Performance & Examination Report: <b>{data['class_name']}</b> ({data['class_code']})", subtitle_style),
        Paragraph(f"Generated: {data['generated_at'].strftime('%d %B %Y, %I:%M %p')} &bull; Arohan Academy Examination Center", meta_style),
    ]
    header_block = create_report_header_flowable(title_paragraphs, logo_size=46, total_width=770)
    if isinstance(header_block, list):
        elements.extend(header_block)
    else:
        elements.append(header_block)
    elements.append(HRFlowable(width="100%", thickness=1, color=colors.HexColor('#CBD5E1'), spaceBefore=5, spaceAfter=7))

    # 2. Executive KPI Table (770 pt total: 154 pt each)
    kpi_data = [
        [
            Paragraph("<b>Enrolled Students</b>", cell_header),
            Paragraph("<b>Scheduled Tests</b>", cell_header),
            Paragraph("<b>Total Submissions</b>", cell_header),
            Paragraph("<b>Class Average</b>", cell_header),
            Paragraph("<b>Pass Rate</b>", cell_header),
        ],
        [
            Paragraph(f"<font size=11><b>{data['total_students']}</b></font>", cell_bold),
            Paragraph(f"<font size=11><b>{data['total_exams_scheduled']}</b></font>", cell_bold),
            Paragraph(f"<font size=11><b>{data['total_attempts_count']}</b></font>", cell_bold),
            Paragraph(f"<font size=11><b>{data['class_avg_pct']}%</b></font>", cell_bold),
            Paragraph(f"<font size=11><b>{data['overall_pass_rate']}%</b></font>", cell_bold),
        ]
    ]

    kpi_table = Table(kpi_data, colWidths=[154, 154, 154, 154, 154])
    kpi_table.setStyle(TableStyle([
        ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#1E3A8A')),
        ('BACKGROUND', (0, 1), (-1, 1), colors.HexColor('#F8FAFC')),
        ('ALIGN', (0, 0), (-1, -1), 'CENTER'),
        ('VALIGN', (0, 0), (-1, -1), 'MIDDLE'),
        ('BOTTOMPADDING', (0, 0), (-1, -1), 5),
        ('TOPPADDING', (0, 0), (-1, -1), 5),
        ('GRID', (0, 0), (-1, -1), 0.5, colors.HexColor('#CBD5E1')),
    ]))
    elements.append(kpi_table)
    elements.append(Spacer(1, 8))

    # 3. Table 1: Exam Schedule & Attempts Breakdown (770 pt total)
    elements.append(Paragraph("1. Mock Tests Schedule vs. Student Attempts", heading_section))

    t1_headers = [
        Paragraph("<b>Mock Test Title</b>", cell_header),
        Paragraph("<b>Marks</b>", cell_header),
        Paragraph("<b>Enrolled</b>", cell_header),
        Paragraph("<b>Attempted</b>", cell_header),
        Paragraph("<b>Pending</b>", cell_header),
        Paragraph("<b>Passed</b>", cell_header),
        Paragraph("<b>Avg %</b>", cell_header),
        Paragraph("<b>Highest</b>", cell_header),
    ]

    t1_rows = [t1_headers]
    for ex in data['exam_stats']:
        t1_rows.append([
            Paragraph(f"<b>{ex['title']}</b><br/><font color='#64748B'>{ex['subject']}</font>", cell_style),
            Paragraph(str(ex['total_marks']), cell_style),
            Paragraph(str(ex['enrolled_students']), cell_style),
            Paragraph(f"<b>{ex['attempted_students']}</b>", cell_style),
            Paragraph(f"<font color='{'#B45309' if ex['pending_students'] > 0 else '#15803D'}'><b>{ex['pending_students']}</b></font>", cell_style),
            Paragraph(str(ex['passed_count']), cell_style),
            Paragraph(f"<b>{ex['avg_percentage']}%</b>", cell_style),
            Paragraph(str(ex['highest_score']), cell_style),
        ])

    t1_table = Table(t1_rows, colWidths=[190, 60, 70, 80, 80, 70, 80, 140])
    t1_table.setStyle(TableStyle([
        ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#0F172A')),
        ('ALIGN', (1, 0), (-1, -1), 'CENTER'),
        ('VALIGN', (0, 0), (-1, -1), 'MIDDLE'),
        ('ROWBACKGROUNDS', (0, 1), (-1, -1), [colors.white, colors.HexColor('#F8FAFC')]),
        ('BOTTOMPADDING', (0, 0), (-1, -1), 4),
        ('TOPPADDING', (0, 0), (-1, -1), 4),
        ('LEFTPADDING', (0, 0), (-1, -1), 6),
        ('RIGHTPADDING', (0, 0), (-1, -1), 6),
        ('GRID', (0, 0), (-1, -1), 0.5, colors.HexColor('#CBD5E1')),
    ]))
    elements.append(t1_table)
    elements.append(Spacer(1, 8))

    # 4. Table 2: Student-wise Breakdown (770 pt total)
    elements.append(Paragraph("2. Enrolled Students Progress Roster", heading_section))

    t2_headers = [
        Paragraph("<b>Roll No</b>", cell_header),
        Paragraph("<b>Student Name</b>", cell_header),
        Paragraph("<b>Assigned</b>", cell_header),
        Paragraph("<b>Attempted</b>", cell_header),
        Paragraph("<b>Pending</b>", cell_header),
        Paragraph("<b>Passed</b>", cell_header),
        Paragraph("<b>Avg %</b>", cell_header),
        Paragraph("<b>Status</b>", cell_header),
    ]

    t2_rows = [t2_headers]
    for st in data['student_stats']:
        t2_rows.append([
            Paragraph(st['roll_number'], cell_bold),
            Paragraph(st['name'], cell_style),
            Paragraph(str(st['assigned_exams']), cell_style),
            Paragraph(f"<b>{st['attempted_exams']}</b>", cell_style),
            Paragraph(f"<font color='{'#B45309' if st['pending_exams'] > 0 else '#15803D'}'><b>{st['pending_exams']}</b></font>", cell_style),
            Paragraph(str(st['passed_count']), cell_style),
            Paragraph(f"<b>{st['avg_percentage']}%</b>", cell_style),
            Paragraph(st['status'], cell_style),
        ])

    t2_table = Table(t2_rows, colWidths=[65, 205, 70, 80, 80, 70, 80, 120])
    t2_table.setStyle(TableStyle([
        ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#1E3A8A')),
        ('ALIGN', (0, 0), (0, -1), 'CENTER'),
        ('ALIGN', (2, 0), (-1, -1), 'CENTER'),
        ('VALIGN', (0, 0), (-1, -1), 'MIDDLE'),
        ('ROWBACKGROUNDS', (0, 1), (-1, -1), [colors.white, colors.HexColor('#F8FAFC')]),
        ('BOTTOMPADDING', (0, 0), (-1, -1), 4),
        ('TOPPADDING', (0, 0), (-1, -1), 4),
        ('LEFTPADDING', (0, 0), (-1, -1), 6),
        ('RIGHTPADDING', (0, 0), (-1, -1), 6),
        ('GRID', (0, 0), (-1, -1), 0.5, colors.HexColor('#CBD5E1')),
    ]))
    elements.append(t2_table)

    doc.build(elements)
    buf.seek(0)
    return buf.getvalue()


def generate_student_pdf(data):
    """
    Generates a print-ready Landscape A4 PDF document for Student scorecard using ReportLab.
    """
    buf = io.BytesIO()
    doc = SimpleDocTemplate(
        buf,
        pagesize=landscape(A4),
        leftMargin=36,
        rightMargin=36,
        topMargin=28,
        bottomMargin=28
    )

    styles = getSampleStyleSheet()

    title_style = ParagraphStyle(
        'DocTitle',
        parent=styles['Heading1'],
        fontName='Helvetica-Bold',
        fontSize=15,
        leading=18,
        textColor=colors.HexColor('#1E3A8A'),
        spaceAfter=2
    )

    subtitle_style = ParagraphStyle(
        'DocSubtitle',
        parent=styles['Normal'],
        fontName='Helvetica-Bold',
        fontSize=10,
        leading=13,
        textColor=colors.HexColor('#D97706'),
        spaceAfter=3
    )

    meta_style = ParagraphStyle(
        'DocMeta',
        parent=styles['Normal'],
        fontName='Helvetica',
        fontSize=8,
        leading=10,
        textColor=colors.HexColor('#64748B'),
        spaceAfter=0
    )

    heading_section = ParagraphStyle(
        'SectionHeading',
        parent=styles['Heading2'],
        fontName='Helvetica-Bold',
        fontSize=11,
        leading=14,
        textColor=colors.HexColor('#0F172A'),
        spaceBefore=8,
        spaceAfter=6
    )

    cell_style = ParagraphStyle(
        'CellText',
        parent=styles['Normal'],
        fontName='Helvetica',
        fontSize=8.5,
        leading=11,
        textColor=colors.HexColor('#1E293B')
    )

    cell_style_center = ParagraphStyle(
        'CellTextCenter',
        parent=styles['Normal'],
        fontName='Helvetica',
        fontSize=8.5,
        leading=11,
        textColor=colors.HexColor('#1E293B'),
        alignment=1
    )

    cell_bold = ParagraphStyle(
        'CellBold',
        parent=styles['Normal'],
        fontName='Helvetica-Bold',
        fontSize=8.5,
        leading=11,
        textColor=colors.HexColor('#0F172A'),
        alignment=1
    )

    cell_header = ParagraphStyle(
        'CellHeader',
        parent=styles['Normal'],
        fontName='Helvetica-Bold',
        fontSize=8.5,
        leading=11,
        textColor=colors.white,
        alignment=1
    )

    elements = []

    # 1. School Header Banner with Logo on Left
    title_paragraphs = [
        Paragraph("Arohan Academy English School", title_style),
        Paragraph(f"OFFICIAL STUDENT SCORECARD: <b>{data['name'].upper()}</b> (ROLL: {data['roll_number']})", subtitle_style),
        Paragraph(f"Academic Year Examination Record &bull; Generated: {data['generated_at'].strftime('%d %B %Y, %I:%M %p')}", meta_style),
    ]
    header_block = create_report_header_flowable(title_paragraphs, logo_size=46, total_width=770)
    if isinstance(header_block, list):
        elements.extend(header_block)
    else:
        elements.append(header_block)
    elements.append(HRFlowable(width="100%", thickness=1, color=colors.HexColor('#CBD5E1'), spaceBefore=5, spaceAfter=7))

    # 2. Prominent Student Profile Card (770 pt total: 110 + 275 + 110 + 275)
    profile_label_style = ParagraphStyle(
        'ProfileLabel',
        parent=styles['Normal'],
        fontName='Helvetica-Bold',
        fontSize=8,
        leading=10,
        textColor=colors.HexColor('#64748B')
    )
    profile_value_style = ParagraphStyle(
        'ProfileValue',
        parent=styles['Normal'],
        fontName='Helvetica-Bold',
        fontSize=9,
        leading=11,
        textColor=colors.HexColor('#0F172A')
    )
    profile_name_style = ParagraphStyle(
        'ProfileName',
        parent=styles['Normal'],
        fontName='Helvetica-Bold',
        fontSize=11,
        leading=13,
        textColor=colors.HexColor('#1E3A8A')
    )

    profile_data = [
        [
            Paragraph("<b>STUDENT NAME:</b>", profile_label_style),
            Paragraph(f"<b>{data['name']}</b>", profile_name_style),
            Paragraph("<b>ROLL NUMBER:</b>", profile_label_style),
            Paragraph(f"<b>{data['roll_number']}</b>", profile_name_style),
        ],
        [
            Paragraph("<b>CLASS / BATCH:</b>", profile_label_style),
            Paragraph(f"<b>{data['class_name']}</b>", profile_value_style),
            Paragraph("<b>CONTACT PHONE:</b>", profile_label_style),
            Paragraph(f"{data['phone']}", profile_value_style),
        ],
        [
            Paragraph("<b>EXAM CENTER:</b>", profile_label_style),
            Paragraph("Arohan Academy Examination Center", profile_value_style),
            Paragraph("<b>EMAIL ADDRESS:</b>", profile_label_style),
            Paragraph(f"{data['email']}", profile_value_style),
        ]
    ]

    profile_table = Table(profile_data, colWidths=[110, 275, 110, 275])
    profile_table.setStyle(TableStyle([
        ('BACKGROUND', (0, 0), (-1, -1), colors.HexColor('#F8FAFC')),
        ('VALIGN', (0, 0), (-1, -1), 'MIDDLE'),
        ('LEFTPADDING', (0, 0), (-1, -1), 8),
        ('RIGHTPADDING', (0, 0), (-1, -1), 8),
        ('TOPPADDING', (0, 0), (-1, -1), 4),
        ('BOTTOMPADDING', (0, 0), (-1, -1), 4),
        ('BOX', (0, 0), (-1, -1), 1, colors.HexColor('#CBD5E1')),
        ('INNERGRID', (0, 0), (-1, -1), 0.5, colors.HexColor('#E2E8F0')),
    ]))
    elements.append(profile_table)
    elements.append(Spacer(1, 6))

    # 3. Student KPI Table (770 pt total: 128*4 + 129*2)
    kpi_data = [
        [
            Paragraph("<b>Assigned Tests</b>", cell_header),
            Paragraph("<b>Attempted Tests</b>", cell_header),
            Paragraph("<b>Pending Tests</b>", cell_header),
            Paragraph("<b>Submissions</b>", cell_header),
            Paragraph("<b>Average %</b>", cell_header),
            Paragraph("<b>Best Score %</b>", cell_header),
        ],
        [
            Paragraph(f"<font size=11><b>{data['total_assigned_exams']}</b></font>", cell_bold),
            Paragraph(f"<font size=11><b>{data['attempted_exams_count']}</b></font>", cell_bold),
            Paragraph(f"<font size=11 color='{'#B45309' if data['pending_exams_count'] > 0 else '#15803D'}'><b>{data['pending_exams_count']}</b></font>", cell_bold),
            Paragraph(f"<font size=11><b>{data['total_attempts_count']}</b></font>", cell_bold),
            Paragraph(f"<font size=11><b>{data['overall_avg_pct']}%</b></font>", cell_bold),
            Paragraph(f"<font size=11 color='#15803D'><b>{data['best_pct']}%</b></font>", cell_bold),
        ]
    ]

    kpi_table = Table(kpi_data, colWidths=[128, 128, 128, 128, 129, 129])
    kpi_table.setStyle(TableStyle([
        ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#1E3A8A')),
        ('BACKGROUND', (0, 1), (-1, 1), colors.HexColor('#F8FAFC')),
        ('ALIGN', (0, 0), (-1, -1), 'CENTER'),
        ('VALIGN', (0, 0), (-1, -1), 'MIDDLE'),
        ('BOTTOMPADDING', (0, 0), (-1, -1), 4),
        ('TOPPADDING', (0, 0), (-1, -1), 4),
        ('GRID', (0, 0), (-1, -1), 0.5, colors.HexColor('#CBD5E1')),
    ]))
    elements.append(kpi_table)
    elements.append(Spacer(1, 8))

    # 4. Table 1: Attempt History (770 pt total: 40 + 185 + 150 + 130 + 65 + 75 + 65 + 60)
    elements.append(Paragraph(f"1. Chronological Attempt History: <b>{data['name']}</b>", heading_section))

    att_headers = [
        Paragraph("<b>#</b>", cell_header),
        Paragraph("<b>Mock Test Title</b>", cell_header),
        Paragraph("<b>Subject</b>", cell_header),
        Paragraph("<b>Date & Time</b>", cell_header),
        Paragraph("<b>Time (min)</b>", cell_header),
        Paragraph("<b>Score / Total</b>", cell_header),
        Paragraph("<b>Percentage</b>", cell_header),
        Paragraph("<b>Result</b>", cell_header),
    ]

    att_rows = [att_headers]
    for row in data['attempt_rows']:
        att_rows.append([
            Paragraph(f"#{row['attempt_number']}", cell_bold),
            Paragraph(f"<b>{row['exam_title']}</b>", cell_style),
            Paragraph(row['subject'], cell_style),
            Paragraph(row['started_at'].strftime('%d %b %Y, %I:%M %p'), cell_style_center),
            Paragraph(f"{row['time_taken_minutes']} min", cell_style_center),
            Paragraph(f"<b>{row['score_obtained']} / {row['max_score']}</b>", cell_style_center),
            Paragraph(f"<b>{row['percentage']}%</b>", cell_style_center),
            Paragraph(
                f"<font color='{'#15803D' if row['is_passed'] else '#B91C1C'}'><b>{'PASS' if row['is_passed'] else 'FAIL'}</b></font>",
                cell_style_center
            ),
        ])

    if len(att_rows) > 1:
        att_table = Table(att_rows, colWidths=[40, 185, 150, 130, 65, 75, 65, 60])
        att_table.setStyle(TableStyle([
            ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#0F172A')),
            ('ALIGN', (0, 0), (-1, 0), 'CENTER'),
            ('ALIGN', (0, 1), (0, -1), 'CENTER'),
            ('ALIGN', (1, 1), (2, -1), 'LEFT'),
            ('ALIGN', (3, 1), (-1, -1), 'CENTER'),
            ('VALIGN', (0, 0), (-1, -1), 'MIDDLE'),
            ('ROWBACKGROUNDS', (0, 1), (-1, -1), [colors.white, colors.HexColor('#F8FAFC')]),
            ('BOTTOMPADDING', (0, 0), (-1, -1), 4),
            ('TOPPADDING', (0, 0), (-1, -1), 4),
            ('LEFTPADDING', (0, 0), (-1, -1), 6),
            ('RIGHTPADDING', (0, 0), (-1, -1), 6),
            ('GRID', (0, 0), (-1, -1), 0.5, colors.HexColor('#CBD5E1')),
        ]))
        elements.append(att_table)
    else:
        elements.append(Paragraph("<i>No attempts recorded yet for this student.</i>", meta_style))

    elements.append(Spacer(1, 8))

    # 5. Table 2: Pending Tests (770 pt total: 230 + 160 + 80 + 90 + 90 + 120)
    if data['pending_exams']:
        elements.append(Paragraph("2. Pending Mock Tests (Yet to Attempt)", heading_section))
        p_headers = [
            Paragraph("<b>Mock Test Title</b>", cell_header),
            Paragraph("<b>Subject</b>", cell_header),
            Paragraph("<b>Duration</b>", cell_header),
            Paragraph("<b>Total Marks</b>", cell_header),
            Paragraph("<b>Pass %</b>", cell_header),
            Paragraph("<b>Status</b>", cell_header),
        ]
        p_rows = [p_headers]
        for p in data['pending_exams']:
            p_rows.append([
                Paragraph(f"<b>{p.title}</b>", cell_style),
                Paragraph(p.subject or 'General', cell_style),
                Paragraph(f"{p.duration_minutes}m", cell_style_center),
                Paragraph(str(p.total_marks), cell_style_center),
                Paragraph(f"{p.passing_percentage}%", cell_style_center),
                Paragraph("<font color='#B45309'><b>PENDING</b></font>", cell_style_center),
            ])
        p_table = Table(p_rows, colWidths=[230, 160, 80, 90, 90, 120])
        p_table.setStyle(TableStyle([
            ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#B45309')),
            ('ALIGN', (0, 0), (-1, 0), 'CENTER'),
            ('ALIGN', (0, 1), (1, -1), 'LEFT'),
            ('ALIGN', (2, 1), (-1, -1), 'CENTER'),
            ('VALIGN', (0, 0), (-1, -1), 'MIDDLE'),
            ('ROWBACKGROUNDS', (0, 1), (-1, -1), [colors.white, colors.HexColor('#FEF3C7')]),
            ('BOTTOMPADDING', (0, 0), (-1, -1), 4),
            ('TOPPADDING', (0, 0), (-1, -1), 4),
            ('LEFTPADDING', (0, 0), (-1, -1), 6),
            ('RIGHTPADDING', (0, 0), (-1, -1), 6),
            ('GRID', (0, 0), (-1, -1), 0.5, colors.HexColor('#CBD5E1')),
        ]))
        elements.append(p_table)

    doc.build(elements)
    buf.seek(0)
    return buf.getvalue()


def generate_attempt_result_pdf(attempt):
    """
    Generates a high-fidelity, print-ready Portrait A4 PDF for a student's
    exam attempt result using ReportLab.
    Includes:
    - Official school header with logo
    - Student candidate profile & test metadata
    - Performance summary & KPI scorecard
    - Detailed question-by-question verification review with color-coded options,
      marks, status badges, and solution explanations.
    """
    from xml.sax.saxutils import escape as xml_escape

    def safe_xml(text):
        if text is None:
            return ""
        return xml_escape(str(text))

    buf = io.BytesIO()
    doc = SimpleDocTemplate(
        buf,
        pagesize=A4,
        leftMargin=32,
        rightMargin=32,
        topMargin=28,
        bottomMargin=28
    )

    styles = getSampleStyleSheet()

    # Typography & Styles
    title_style = ParagraphStyle(
        'AttemptTitle',
        parent=styles['Heading1'],
        fontName='Helvetica-Bold',
        fontSize=15,
        leading=18,
        textColor=colors.HexColor('#1E3A8A'),
        spaceAfter=2
    )
    subtitle_style = ParagraphStyle(
        'AttemptSubtitle',
        parent=styles['Normal'],
        fontName='Helvetica-Bold',
        fontSize=9.5,
        leading=13,
        textColor=colors.HexColor('#B45309'),
        spaceAfter=2
    )
    meta_style = ParagraphStyle(
        'AttemptMeta',
        parent=styles['Normal'],
        fontName='Helvetica',
        fontSize=8,
        leading=10,
        textColor=colors.HexColor('#64748B'),
        spaceAfter=0
    )
    heading_section = ParagraphStyle(
        'SectionHeading',
        parent=styles['Heading2'],
        fontName='Helvetica-Bold',
        fontSize=11,
        leading=14,
        textColor=colors.HexColor('#0F172A'),
        spaceBefore=10,
        spaceAfter=6
    )
    cell_header = ParagraphStyle(
        'CellHeader',
        parent=styles['Normal'],
        fontName='Helvetica-Bold',
        fontSize=8,
        leading=10,
        textColor=colors.white,
        alignment=1
    )
    cell_bold = ParagraphStyle(
        'CellBold',
        parent=styles['Normal'],
        fontName='Helvetica-Bold',
        fontSize=8.5,
        leading=11,
        textColor=colors.HexColor('#0F172A'),
        alignment=1
    )
    cell_text = ParagraphStyle(
        'CellText',
        parent=styles['Normal'],
        fontName='Helvetica',
        fontSize=8.5,
        leading=11,
        textColor=colors.HexColor('#1E293B')
    )
    cell_label = ParagraphStyle(
        'CellLabel',
        parent=styles['Normal'],
        fontName='Helvetica-Bold',
        fontSize=8,
        leading=10,
        textColor=colors.HexColor('#64748B')
    )
    q_title_style = ParagraphStyle(
        'QTitle',
        parent=styles['Normal'],
        fontName='Helvetica-Bold',
        fontSize=9,
        leading=12,
        textColor=colors.HexColor('#0F172A')
    )
    q_badge_style = ParagraphStyle(
        'QBadge',
        parent=styles['Normal'],
        fontName='Helvetica-Bold',
        fontSize=8,
        leading=10,
        alignment=2
    )
    opt_style = ParagraphStyle(
        'OptStyle',
        parent=styles['Normal'],
        fontName='Helvetica',
        fontSize=8.5,
        leading=11,
        textColor=colors.HexColor('#1E293B')
    )
    exp_style = ParagraphStyle(
        'ExpStyle',
        parent=styles['Normal'],
        fontName='Helvetica',
        fontSize=8,
        leading=11,
        textColor=colors.HexColor('#0369A1')
    )

    elements = []
    total_width = 531  # 595 - 64

    # 1. School Header Banner
    exam = attempt.exam
    student = attempt.student
    comp_date_str = attempt.completed_at.strftime('%d %B %Y, %I:%M %p') if attempt.completed_at else 'In Progress'

    title_paragraphs = [
        Paragraph("Arohan Academy English School", title_style),
        Paragraph(f"STUDENT EXAMINATION RESULT &bull; {safe_xml(exam.title).upper()}", subtitle_style),
        Paragraph(f"Official Verification Report &bull; Attempt #{attempt.attempt_number} &bull; {comp_date_str}", meta_style),
    ]
    header_block = create_report_header_flowable(title_paragraphs, logo_size=46, total_width=total_width)
    if isinstance(header_block, list):
        elements.extend(header_block)
    else:
        elements.append(header_block)

    elements.append(HRFlowable(width="100%", thickness=1, color=colors.HexColor('#CBD5E1'), spaceBefore=5, spaceAfter=8))

    # 2. Student & Exam Profile Table
    profile_data = [
        [
            Paragraph("<b>STUDENT NAME:</b>", cell_label),
            Paragraph(f"<b>{safe_xml(student.name)}</b>", cell_bold),
            Paragraph("<b>MOCK TEST:</b>", cell_label),
            Paragraph(f"<b>{safe_xml(exam.title)}</b>", cell_bold),
        ],
        [
            Paragraph("<b>ROLL NUMBER:</b>", cell_label),
            Paragraph(f"<b>{safe_xml(student.roll_number)}</b>", cell_text),
            Paragraph("<b>SUBJECT:</b>", cell_label),
            Paragraph(f"<b>{safe_xml(exam.subject or 'General')}</b>", cell_text),
        ],
        [
            Paragraph("<b>CLASS / BATCH:</b>", cell_label),
            Paragraph(f"<b>{safe_xml(student.student_class.name if student.student_class else 'N/A')}</b>", cell_text),
            Paragraph("<b>PASS MARK %:</b>", cell_label),
            Paragraph(f"<b>{exam.passing_percentage}%</b>", cell_text),
        ],
        [
            Paragraph("<b>TIME TAKEN:</b>", cell_label),
            Paragraph(f"<b>{attempt.formatted_time_taken}</b>", cell_text),
            Paragraph("<b>EXAM CENTER:</b>", cell_label),
            Paragraph("<b>Arohan Academy Examination Portal</b>", cell_text),
        ]
    ]
    profile_table = Table(profile_data, colWidths=[95, 170, 95, 171])
    profile_table.setStyle(TableStyle([
        ('BACKGROUND', (0, 0), (-1, -1), colors.HexColor('#F8FAFC')),
        ('VALIGN', (0, 0), (-1, -1), 'MIDDLE'),
        ('LEFTPADDING', (0, 0), (-1, -1), 7),
        ('RIGHTPADDING', (0, 0), (-1, -1), 7),
        ('TOPPADDING', (0, 0), (-1, -1), 3.5),
        ('BOTTOMPADDING', (0, 0), (-1, -1), 3.5),
        ('BOX', (0, 0), (-1, -1), 1, colors.HexColor('#CBD5E1')),
        ('INNERGRID', (0, 0), (-1, -1), 0.5, colors.HexColor('#E2E8F0')),
    ]))
    elements.append(profile_table)
    elements.append(Spacer(1, 6))

    # 3. Executive KPI Scorecard
    result_text = "PASSED" if attempt.is_passed else "FAILED"
    result_bg = colors.HexColor('#DCFCE7') if attempt.is_passed else colors.HexColor('#FEE2E2')
    result_fg = '#15803D' if attempt.is_passed else '#B91C1C'

    kpi_data = [
        [
            Paragraph("<b>Questions</b>", cell_header),
            Paragraph("<b>Correct</b>", cell_header),
            Paragraph("<b>Incorrect</b>", cell_header),
            Paragraph("<b>Skipped</b>", cell_header),
            Paragraph("<b>Score Obtained</b>", cell_header),
            Paragraph("<b>Final Status</b>", cell_header),
        ],
        [
            Paragraph(f"<font size=11><b>{attempt.total_questions}</b></font>", cell_bold),
            Paragraph(f"<font size=11 color='#15803D'><b>{attempt.correct_answers}</b></font>", cell_bold),
            Paragraph(f"<font size=11 color='#B91C1C'><b>{attempt.wrong_answers}</b></font>", cell_bold),
            Paragraph(f"<font size=11 color='#64748B'><b>{attempt.unattempted_answers}</b></font>", cell_bold),
            Paragraph(f"<font size=10><b>{attempt.score_obtained} / {attempt.max_score}</b><br/>({attempt.percentage}%)</font>", cell_bold),
            Paragraph(f"<font size=10 color='{result_fg}'><b>{result_text}</b><br/>{attempt.grade} ({attempt.grade_title})</font>", cell_bold),
        ]
    ]
    kpi_table = Table(kpi_data, colWidths=[80, 80, 80, 80, 105, 106])
    kpi_table.setStyle(TableStyle([
        ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#1E3A8A')),
        ('BACKGROUND', (0, 1), (-2, 1), colors.HexColor('#F8FAFC')),
        ('BACKGROUND', (5, 1), (5, 1), result_bg),
        ('ALIGN', (0, 0), (-1, -1), 'CENTER'),
        ('VALIGN', (0, 0), (-1, -1), 'MIDDLE'),
        ('BOTTOMPADDING', (0, 0), (-1, -1), 4),
        ('TOPPADDING', (0, 0), (-1, -1), 4),
        ('GRID', (0, 0), (-1, -1), 0.5, colors.HexColor('#CBD5E1')),
    ]))
    elements.append(kpi_table)
    elements.append(Spacer(1, 10))

    # 4. Detailed Question-by-Question Review
    elements.append(Paragraph("Detailed Question-by-Question Verification Review", heading_section))

    answers = attempt.answers.select_related('question').order_by('question__order', 'question__id')

    for idx, ans in enumerate(answers, start=1):
        q = ans.question
        q_elements = []

        # Question Header Box
        if ans.is_correct:
            badge_html = f"<font color='#15803D'><b>[✓] CORRECT (+{q.marks})</b></font>"
            bar_border = colors.HexColor('#16A34A')
        elif ans.selected_option:
            neg_str = f"-{q.negative_marks}" if q.negative_marks > 0 else "0.0"
            badge_html = f"<font color='#B91C1C'><b>[✗] INCORRECT ({neg_str})</b></font>"
            bar_border = colors.HexColor('#DC2626')
        else:
            badge_html = "<font color='#64748B'><b>[-] SKIPPED (0.0)</b></font>"
            bar_border = colors.HexColor('#94A3B8')

        q_stmt = f"<b>Q{q.order or idx}.</b> {safe_xml(q.question_text).replace(chr(10), '<br/>')}"
        q_bar_table = Table(
            [[Paragraph(q_stmt, q_title_style), Paragraph(badge_html, q_badge_style)]],
            colWidths=[415, 115]
        )
        q_bar_table.setStyle(TableStyle([
            ('BACKGROUND', (0, 0), (-1, -1), colors.HexColor('#F1F5F9')),
            ('VALIGN', (0, 0), (-1, -1), 'TOP'),
            ('TOPPADDING', (0, 0), (-1, -1), 5),
            ('BOTTOMPADDING', (0, 0), (-1, -1), 5),
            ('LEFTPADDING', (0, 0), (-1, -1), 7),
            ('RIGHTPADDING', (0, 0), (-1, -1), 7),
            ('BOX', (0, 0), (-1, -1), 1, bar_border),
        ]))
        q_elements.append(q_bar_table)

        # 4 Options
        options_map = [
            ('A', q.option_a),
            ('B', q.option_b),
            ('C', q.option_c),
            ('D', q.option_d),
        ]
        opt_table_rows = []
        opt_table_styles = [
            ('VALIGN', (0, 0), (-1, -1), 'MIDDLE'),
            ('TOPPADDING', (0, 0), (-1, -1), 3),
            ('BOTTOMPADDING', (0, 0), (-1, -1), 3),
            ('LEFTPADDING', (0, 0), (-1, -1), 8),
            ('RIGHTPADDING', (0, 0), (-1, -1), 8),
            ('INNERGRID', (0, 0), (-1, -1), 0.5, colors.HexColor('#E2E8F0')),
            ('BOX', (0, 0), (-1, -1), 0.5, colors.HexColor('#CBD5E1')),
        ]

        for r_idx, (letter, text) in enumerate(options_map):
            is_correct_opt = (q.correct_option == letter)
            is_student_opt = (ans.selected_option == letter)

            clean_text = safe_xml(text)
            if is_correct_opt and is_student_opt:
                opt_html = f"<b>{letter}) {clean_text}</b>"
                tag_html = "<font color='#15803D'><b>[✓ Your Answer &bull; Correct]</b></font>"
                bg_color = colors.HexColor('#DCFCE7')
            elif is_correct_opt:
                opt_html = f"<b>{letter}) {clean_text}</b>"
                tag_html = "<font color='#15803D'><b>[✓ Correct Answer]</b></font>"
                bg_color = colors.HexColor('#ECFDF5')
            elif is_student_opt:
                opt_html = f"<b>{letter}) {clean_text}</b>"
                tag_html = "<font color='#B91C1C'><b>[✗ Your Answer &bull; Incorrect]</b></font>"
                bg_color = colors.HexColor('#FEE2E2')
            else:
                opt_html = f"<b>{letter})</b> {clean_text}"
                tag_html = ""
                bg_color = colors.white

            opt_table_rows.append([
                Paragraph(opt_html, opt_style),
                Paragraph(tag_html, q_badge_style)
            ])
            opt_table_styles.append(('BACKGROUND', (0, r_idx), (-1, r_idx), bg_color))

        opt_table = Table(opt_table_rows, colWidths=[380, 150])
        opt_table.setStyle(TableStyle(opt_table_styles))
        q_elements.append(opt_table)

        # Explanation (if present)
        if q.explanation and q.explanation.strip():
            exp_text = f"<b>💡 Explanation & Solution:</b> {safe_xml(q.explanation.strip()).replace(chr(10), '<br/>')}"
            exp_table = Table([[Paragraph(exp_text, exp_style)]], colWidths=[530])
            exp_table.setStyle(TableStyle([
                ('BACKGROUND', (0, 0), (-1, -1), colors.HexColor('#F0F9FF')),
                ('BOX', (0, 0), (-1, -1), 0.5, colors.HexColor('#BAE6FD')),
                ('LEFTPADDING', (0, 0), (-1, -1), 8),
                ('RIGHTPADDING', (0, 0), (-1, -1), 8),
                ('TOPPADDING', (0, 0), (-1, -1), 4),
                ('BOTTOMPADDING', (0, 0), (-1, -1), 4),
            ]))
            q_elements.append(exp_table)

        q_elements.append(Spacer(1, 8))
        elements.append(KeepTogether(q_elements))

    # Footer note
    elements.append(Spacer(1, 8))
    elements.append(HRFlowable(width="100%", thickness=0.5, color=colors.HexColor('#CBD5E1'), spaceBefore=4, spaceAfter=4))
    elements.append(Paragraph(
        "&copy; 2026 Arohan Academy English School &bull; Automated Online Examination Engine &bull; Official Student Verification Record",
        meta_style
    ))

    doc.build(elements)
    buf.seek(0)
    return buf.getvalue()
